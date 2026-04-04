"""
Excel exporter for scan results with advanced formatting.
"""

from __future__ import annotations

import io
from typing import Tuple

from .base import BaseExporter


class ExcelExporter(BaseExporter):
    """Export scan results to Excel format with advanced formatting."""
    
    def export(self) -> Tuple[bytes, str, str]:
        """Export scan data to Excel format."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
            from openpyxl.cell import WriteOnlyCell
        except ImportError:
            raise ImportError("openpyxl not installed. Install with: pip install openpyxl")
        
        # Create workbook
        wb = openpyxl.Workbook(write_only=True)
        
        # Create summary sheet
        self._create_summary_sheet(wb)
        
        # Create files sheet
        self._create_files_sheet(wb)
        
        # Create categories sheet
        self._create_categories_sheet(wb)
        
        # Save to bytes
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        
        filename = self._generate_filename("xlsx")
        return buf.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename
    
    def _create_summary_sheet(self, wb):
        """Create summary sheet with scan metadata."""
        ws = wb.create_sheet("Resumen")
        
        # Define styles
        header_font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
        cell_font = Font(name='Calibri', size=11)
        
        # Add title
        title_cell = WriteOnlyCell(ws, value="Smart File Organizer - Resumen de Escaneo")
        title_cell.font = Font(name='Calibri', size=16, bold=True, color='0F172A')
        ws.append([title_cell])
        ws.append([])  # Empty row
        
        # Add metadata
        metadata_items = [
            ("ID del Escaneo", self.meta["scan_id"]),
            ("Ruta Escaneada", self.meta["root_path"]),
            ("Fecha del Escaneo", self.meta["scanned_at"]),
            ("Total de Archivos", str(self.meta["total_files"])),
            ("Tamaño Total", self._fmt_size(self.meta["total_size"])),
            ("Duración (segundos)", str(self.meta["duration_sec"]))
        ]
        
        for label, value in metadata_items:
            label_cell = WriteOnlyCell(ws, value=label)
            label_cell.font = header_font
            label_cell.fill = header_fill
            value_cell = WriteOnlyCell(ws, value=value)
            value_cell.font = cell_font
            ws.append([label_cell, value_cell])
    
    def _create_files_sheet(self, wb):
        """Create files sheet with all scan results."""
        ws = wb.create_sheet("Archivos")
        
        # Define styles
        header_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Add headers
        headers = ["Nombre", "Ruta", "Tamaño (bytes)", "Tamaño", "Extensión", "Categoría", "Modificado"]
        header_cells = []
        for header in headers:
            cell = WriteOnlyCell(ws, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            header_cells.append(cell)
        ws.append(header_cells)
        
        # Add file data
        for file_data in self._iter_files():
            row_data = [
                file_data["name"],
                file_data["path"],
                file_data["size"],
                self._fmt_size(file_data["size"]),
                file_data["extension"],
                file_data["category"],
                file_data["modified"]
            ]
            ws.append(row_data)
    
    def _create_categories_sheet(self, wb):
        """Create categories sheet with statistics."""
        ws = wb.create_sheet("Categorías")
        
        # Define styles
        header_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='3B82F6', end_color='3B82F6', fill_type='solid')
        
        # Add headers
        headers = ["Categoría", "Cantidad de Archivos", "Tamaño Total", "Porcentaje"]
        header_cells = []
        for header in headers:
            cell = WriteOnlyCell(ws, value=header)
            cell.font = header_font
            cell.fill = header_fill
            header_cells.append(cell)
        ws.append(header_cells)
        
        # Add category data
        for category in self._get_categories():
            row_data = [
                category["category"],
                category["file_count"],
                self._fmt_size(category["total_size"]),
                f"{category['percentage']}%"
            ]
            ws.append(row_data)
